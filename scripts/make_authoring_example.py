"""Write the example document the authoring format describes.

`docs/AUTHORING_TEMPLATE.md` tells the editor what a document looks like. This
writes one, with every task type in it, so the description has something to
point at and the reader has something to be tested against.

    python scripts/make_authoring_example.py <файл.docx>
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument


HEADER = (
    ("Предмет", "Биология"),
    ("Экзамен", "ЕГЭ"),
    ("Класс", "11"),
    ("Сезон", "21-22"),
    ("Тема", "Цитология"),
)


def _task(document: DocxDocument, number: int, lines: list[str | list[str]]) -> None:
    document.add_paragraph(f"Задание {number}")
    for line in lines:
        if isinstance(line, list):
            table = document.add_table(rows=len(line), cols=2)
            for row, text in zip(table.rows, line):
                left, right = text.split(" | ")
                row.cells[0].text = left
                row.cells[1].text = right
            continue
        document.add_paragraph(line)


def build(path: Path) -> None:
    document = Document()
    for name, value in HEADER:
        document.add_paragraph(f"{name}: {value}")

    _task(document, 1, [
        "Тип: один ответ",
        "Условие:",
        "В какой фазе митоза хромосомы выстраиваются по экватору клетки?",
        "Варианты:",
        "1) Профаза",
        "2) Метафаза",
        "3) Анафаза",
        "4) Телофаза",
        "Ответ: 2",
        "Решение:",
        "В метафазе хромосомы выстраиваются в экваториальной плоскости.",
    ])

    _task(document, 2, [
        "Тип: несколько ответов",
        "Условие:",
        "Выберите органоиды, окружённые двумя мембранами.",
        "Варианты:",
        "1) Митохондрия",
        "2) Рибосома",
        "3) Пластида",
        "4) Лизосома",
        "Ответ: 1, 3",
    ])

    _task(document, 3, [
        "Тип: соответствие",
        "Тема: Методы биологических исследований",
        "Условие:",
        "Установите соответствие между методом и тем, что он изучает.",
        "Пункты:",
        "А) Цитогенетический",
        "Б) Популяционно-статистический",
        "В) Гибридологический",
        "Варианты:",
        "1) Кариотип в метафазной клетке",
        "2) Распространение гена в популяции",
        "3) Наследование признака при скрещивании",
        "4) Состав белков клетки",
        "Ответ: 123",
    ])

    _task(document, 4, [
        "Тип: последовательность",
        "Условие:",
        "Расставьте стадии развития в порядке их наступления.",
        "Пункты:",
        "А) Первая стадия",
        "Б) Вторая стадия",
        "В) Третья стадия",
        "Варианты:",
        "1) Зигота",
        "2) Бластула",
        "3) Гаструла",
        "Ответ: 123",
    ])

    _task(document, 5, [
        "Тип: число",
        "Условие:",
        "Сколько хромосом содержит соматическая клетка человека?",
        "Ответ: 46",
    ])

    _task(document, 6, [
        "Тип: короткий ответ",
        "Условие:",
        "Рассмотрите таблицу и назовите метод, пропущенный в первой строке.",
        [
            "Метод | Применение метода",
            "___ | Изучение кариотипа в метафазной клетке под микроскопом",
            "Популяционно-статистический | Изучение распространения гена в популяции",
        ],
        "Ответ:",
        "цитогенетический",
        "кариотипирование",
        "цитологический",
    ])

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path, help="Куда записать образец")
    arguments = parser.parse_args(argv)
    build(arguments.output)
    print(f"OK {arguments.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
