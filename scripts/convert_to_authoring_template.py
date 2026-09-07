"""Rewrite editorial documents in the authoring format.

Everything the current converter reads without guessing is filled in. Everything
it would have to guess is left empty and marked, so the guess belongs to the
editor and not to a program.

    python scripts/convert_to_authoring_template.py <исходный.docx> <результат.docx>
    python scripts/convert_to_authoring_template.py <папка> <папка> --report <отчёт.md>
"""

from __future__ import annotations

import argparse
import io
from pathlib import Path
import re
import sys

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Cm

SCRIPTS_ROOT = Path(__file__).resolve().parent
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

import authoring_document as authoring  # noqa: E402
import import_sharepoint_diagnostics as legacy  # noqa: E402

REVIEW = "ТРЕБУЕТ ПРОВЕРКИ"
NUMBERED_LINE = re.compile(r"^\s*(\d+)[.)]\s+(\S.*)$")
SKIP_REASONS = {
    "irregular_key": "ключ не удалось прочитать",
    "glued_answer": "в ключе пропал пробел между словами",
    "unreadable_matching": "таблицу пар не удалось прочитать",
    "missing_figure": "условие ссылается на рисунок, которого нет",
    "external_resource": "задание ссылается на внешний файл или запись",
    "open_answer": "ответ развёрнутый, машинной проверке не поддаётся",
    "prompt_too_long": "условие длиннее допустимого",
    "empty_prompt": "условие пустое",
}
KIND_NAMES = {
    "single": authoring.SINGLE,
    "multiple": authoring.MULTIPLE,
    "matching": authoring.MATCHING,
    "text": authoring.SHORT,
}
COLUMNS = (
    authoring.SINGLE, authoring.MULTIPLE, authoring.MATCHING,
    authoring.SEQUENCE, authoring.NUMBER, authoring.SHORT,
)


def _mark(reason: str) -> str:
    return f"{REVIEW} ({reason})"


def _header(document: DocxDocument, path: Path) -> list[str]:
    """Fill the header from the file name, and say which fields it could not."""
    match = legacy.FILENAME.match(path.stem)
    if match is None:
        values = {name: _mark("имя файла не разобрано") for name in authoring.HEADER_FIELDS}
    else:
        subject = legacy.SUBJECT_CODES.get(match.group("subject"))
        values = {
            "Предмет": legacy.SUBJECT_NAMES.get(subject, _mark("предмет не распознан")),
            "Экзамен": match.group("exam"),
            "Класс": "11" if match.group("exam") == "ЕГЭ" else "9",
            "Сезон": f"{match.group('start')}-{match.group('end')}",
            "Тема": _mark("тема в имени файла не указана"),
        }
    for name in authoring.HEADER_FIELDS:
        document.add_paragraph(f"{name}: {values[name]}")
    return [name for name, value in values.items() if value.startswith(REVIEW)]


def _numbered_run(blocks: list[str]) -> tuple[list[str], list[str]]:
    """Split the list of choices out of the condition, wherever it sits.

    A sequence task prints its choices inside the condition, sometimes at the
    end and sometimes above the text they belong to. A run of lines numbered
    1, 2, 3 without a gap is that list and nothing else, so no guess is needed.
    The rest of the condition keeps its order.
    """
    best: tuple[int, int] = (0, 0)
    start = None
    for index, line in enumerate([*blocks, ""]):
        found = NUMBERED_LINE.match(line) if line else None
        if found is not None and int(found.group(1)) == (index - start + 1 if start is not None else 1):
            start = index if start is None else start
            continue
        if start is not None and index - start > best[1] - best[0]:
            best = (start, index)
        start = index if found is not None and found.group(1) == "1" else None
    first, last = best
    if last - first < 2:
        return blocks, []
    options = [NUMBERED_LINE.match(line).group(2).strip() for line in blocks[first:last]]
    return blocks[:first] + blocks[last:], options


def _write_prompt(document: DocxDocument, blocks: list[str], task: legacy.SourceTask) -> None:
    document.add_paragraph("Условие:")
    if task.prompt_nodes:
        for node in task.prompt_nodes:
            if isinstance(node, legacy.PromptParagraph):
                text = legacy.strip_answer_sheet_instructions(node.text)
                if text:
                    document.add_paragraph(text)
                for payload in node.images:
                    try:
                        document.add_picture(io.BytesIO(payload), width=Cm(12))
                    except Exception:  # noqa: BLE001 - the editor can replace a figure
                        document.add_paragraph(_mark("рисунок не удалось перенести"))
            else:
                table = node.table
                columns = max(table.columns, 1)
                written = document.add_table(rows=len(table.rows), cols=columns)
                for row, source_row in zip(written.rows, table.rows):
                    for cell, lines in zip(row.cells, source_row):
                        cell.text = "\n".join(lines)
        return
    for block in blocks:
        document.add_paragraph(block)
    for table in task.prompt_tables:
        columns = max(table.columns, 1)
        written = document.add_table(rows=len(table.rows), cols=columns)
        for row, source_row in zip(written.rows, table.rows):
            for cell, lines in zip(row.cells, source_row):
                cell.text = "\n".join(lines)
    for payload in task.images:
        try:
            document.add_picture(io.BytesIO(payload), width=Cm(12))
        except Exception:  # noqa: BLE001 - an unreadable figure is the editor's call
            document.add_paragraph(_mark("рисунок не удалось перенести"))


def _condition(task: legacy.SourceTask) -> list[str]:
    """The condition without the sentences that describe a paper answer sheet."""
    text = legacy.strip_answer_sheet_instructions("\n".join(task.prompt_blocks))
    return [line for line in text.split("\n") if line]


def convert(source: Path, target: Path) -> dict[str, object]:
    tasks = legacy.parse_document(source)
    document = Document()
    marked_header = _header(document, source)
    counts = dict.fromkeys(COLUMNS, 0)
    marked: list[tuple[int, str]] = []
    renumbered: list[tuple[int, int]] = []

    for position, task in enumerate(tasks, start=1):
        if task.number != position:
            renumbered.append((task.number, position))
        document.add_paragraph(f"Задание {position}")
        kind, payload = legacy.classify(task)
        blocks = _condition(task)

        def note(reason: str) -> None:
            marked.append((position, reason))

        if kind == "skip":
            reason = SKIP_REASONS.get(str(payload), str(payload))
            document.add_paragraph(f"Тип: {_mark(reason)}")
            _write_prompt(document, blocks, task)
            document.add_paragraph(f"Ответ: {_mark(reason)}")
            note(reason)
        elif kind in {"single", "multiple"}:
            if kind == "single":
                if not isinstance(payload, legacy.SingleAnswerSpec):
                    raise TypeError(
                        f"Unexpected answer spec for single: {type(payload).__name__}"
                    )
                indices = payload.indices
            elif isinstance(payload, legacy.MultipleAnswerSpec):
                indices = payload.indices
            else:
                raise TypeError(f"Unexpected answer spec for {kind}: {type(payload).__name__}")
            document.add_paragraph(f"Тип: {KIND_NAMES[kind]}")
            _write_prompt(document, blocks, task)
            document.add_paragraph("Варианты:")
            for index, label in enumerate(task.options, start=1):
                document.add_paragraph(f"{index}) {label}")
            document.add_paragraph(
                "Ответ: " + ", ".join(str(index) for index in indices)
            )
            counts[KIND_NAMES[kind]] += 1
        elif kind == "matching":
            if not isinstance(payload, legacy.MatchingAnswerSpec):
                raise TypeError(f"Unexpected answer spec for matching: {type(payload).__name__}")
            document.add_paragraph(f"Тип: {authoring.MATCHING}")
            _write_prompt(document, blocks, task)
            document.add_paragraph("Пункты:")
            for index, item in enumerate(payload.items):
                document.add_paragraph(f"{chr(ord('А') + index)}) {legacy.clean_line(item)}")
            document.add_paragraph("Варианты:")
            for digit, label in payload.options:
                document.add_paragraph(f"{digit}) {legacy.clean_line(label)}")
            document.add_paragraph(f"Ответ: {payload.key}")
            counts[authoring.MATCHING] += 1
        elif kind == "text":
            if not isinstance(payload, legacy.TextAnswerSpec):
                raise TypeError(f"Unexpected answer spec for text: {type(payload).__name__}")
            document.add_paragraph(f"Тип: {authoring.SHORT}")
            _write_prompt(document, blocks, task)
            document.add_paragraph("Ответ:")
            for wording in payload.correct:
                document.add_paragraph(wording)
            counts[authoring.SHORT] += 1
        else:  # input
            if not isinstance(payload, legacy.InputAnswerSpec):
                raise TypeError(f"Unexpected answer spec for input: {type(payload).__name__}")
            values = payload.correct
            key = values[0]
            if key.isdigit() and len(key) > 1:
                rest, options = _numbered_run(blocks)
                usable = options and all(1 <= int(digit) <= len(options) for digit in key)
                document.add_paragraph(f"Тип: {authoring.SEQUENCE}")
                _write_prompt(document, rest if usable else blocks, task)
                if usable:
                    document.add_paragraph("Варианты:")
                    for index, label in enumerate(options, start=1):
                        document.add_paragraph(f"{index}) {label}")
                    counts[authoring.SEQUENCE] += 1
                else:
                    reason = "перенесите список вариантов из условия в поле «Варианты»"
                    document.add_paragraph(f"Варианты: {_mark(reason)}")
                    note(reason)
                document.add_paragraph(f"Ответ: {key}")
            else:
                document.add_paragraph(f"Тип: {authoring.NUMBER}")
                _write_prompt(document, blocks, task)
                document.add_paragraph(f"Ответ: {key}")
                counts[authoring.NUMBER] += 1

        if task.solution:
            document.add_paragraph("Решение:")
            for line in task.solution:
                document.add_paragraph(line)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return {
        "source": source.name,
        "target": target.name,
        "tasks": len(tasks),
        "counts": counts,
        "marked": marked,
        "header": marked_header,
        "renumbered": renumbered,
    }


def write_report(path: Path, reports: list[dict[str, object]]) -> None:
    lines = [
        "# Перевод редакционных документов в авторский формат",
        "",
        "Сгенерировано `python scripts/convert_to_authoring_template.py <папка> <папка>`.",
        "",
        "Тип задания взят тем же правилом, что и у импортёра каталога. Там, где",
        "правило не даёт однозначного ответа, поле помечено словами «ТРЕБУЕТ",
        "ПРОВЕРКИ» и причиной. Строки про запись ответа в бланк удалены: бланка в",
        "приложении нет. Задания перенумерованы подряд, потому что авторский",
        "документ ведёт свою нумерацию.",
        "",
        "## Итоги",
        "",
        "| Документ | Заданий | " + " | ".join(COLUMNS) + " | Ждут методиста |",
        "|---" * (len(COLUMNS) + 3) + "|",
    ]
    total_tasks = total_marked = 0
    for report in reports:
        counts = report["counts"]
        marked = len(report["marked"]) + len(report["header"])
        total_tasks += int(report["tasks"])
        total_marked += marked
        lines.append(
            f"| {report['source']} | {report['tasks']} | "
            + " | ".join(str(counts[name]) for name in COLUMNS)
            + f" | {marked} |"
        )
    lines.extend([
        "",
        f"Всего документов {len(reports)}, заданий {total_tasks}, "
        f"полей на проверку методисту {total_marked}.",
        "",
        "## Что именно ждёт методиста",
        "",
        "| Документ | Задание | Причина |",
        "|---|---|---|",
    ])
    for report in reports:
        for name in report["header"]:
            lines.append(f"| {report['source']} | шапка | поле «{name}» не выводится из имени файла |")
        for number, reason in report["marked"]:
            lines.append(f"| {report['source']} | {number} | {reason} |")
    renumbered = [report for report in reports if report["renumbered"]]
    if renumbered:
        lines.extend(["", "## Перенумерованные задания", ""])
        for report in renumbered:
            pairs = ", ".join(f"{was} → {now}" for was, now in report["renumbered"])
            lines.append(f"- {report['source']}: {pairs}")
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Редакционный документ или папка")
    parser.add_argument("target", type=Path, help="Куда записать результат")
    parser.add_argument("--report", type=Path, help="Куда записать отчёт по папке")
    arguments = parser.parse_args(argv)

    if arguments.source.is_dir():
        sources = sorted(arguments.source.rglob("*.docx"))
        sources = [path for path in sources if not path.name.startswith("~$")]
        reports = [convert(path, arguments.target / path.name) for path in sources]
        if arguments.report is not None:
            write_report(arguments.report, reports)
        waiting = sum(len(r["marked"]) + len(r["header"]) for r in reports)
        print(f"документов {len(reports)}, полей на проверку методисту {waiting}")
        return 0

    report = convert(arguments.source, arguments.target)
    print(f"{arguments.target}")
    print(f"  заданий: {report['tasks']}, ждут методиста: "
          f"{len(report['marked']) + len(report['header'])}")
    for number, reason in report["marked"]:
        print(f"  задание {number}: {reason}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
