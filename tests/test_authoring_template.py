import io
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from authoring_document import NEEDS_REVIEW, check_document, read_document  # noqa: E402
import convert_to_authoring_template as converter  # noqa: E402


def _png() -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), (10, 120, 200)).save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _source(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Задание 1")
    prompt = document.add_paragraph("Чему равна масса H")
    prompt.add_run("2").font.subscript = True
    prompt.add_run("O?")
    document.add_paragraph("Варианты:")
    document.add_paragraph("16 г")
    document.add_paragraph("18 г")
    document.add_paragraph("Ответ:")
    document.add_paragraph("2")

    document.add_paragraph("Задание 2")
    document.add_paragraph("Выберите верные суждения.")
    document.add_paragraph("Варианты:")
    for label in ("первое", "второе", "третье"):
        document.add_paragraph(label)
    document.add_paragraph("Решение:")
    document.add_paragraph("Верны первое и третье.")
    document.add_paragraph("Ответ:")
    document.add_paragraph("1#3")

    document.add_paragraph("Задание 3")
    document.add_paragraph("Установите соответствие между событиями и годами.")
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "События"
    table.cell(0, 1).text = "Годы"
    table.cell(1, 0).text = "А. Восстание декабристов"
    table.cell(1, 1).text = "1. 1185"
    table.cell(2, 0).text = "Б. Манифест о вольности"
    table.cell(2, 1).text = "2. 1825"
    table.cell(3, 0).text = "В. Поход Игоря"
    table.cell(3, 1).text = "3. 1762"
    document.add_paragraph("В ответ запишите последовательность цифр, соответствующую буквам АБВ.")
    document.add_paragraph("Ответ:")
    document.add_paragraph("231")

    document.add_paragraph("Задание 4")
    document.add_paragraph("Найдите значение выражения на рисунке.")
    document.add_picture(_png(), width=Inches(1))
    document.add_paragraph("Ход решения")
    document.add_paragraph("Считаем.")
    document.add_paragraph("Ответ")
    document.add_paragraph("1,5")

    document.add_paragraph("Задание 5")
    document.add_paragraph("Назовите метод.")
    document.add_paragraph("Ответ:")
    document.add_paragraph("цитогенетический#кариотипирование")

    document.add_paragraph("Задание 6")
    document.add_paragraph("Раскройте понятие «деньги».")
    document.add_paragraph("Решение:")
    document.add_paragraph("Любой развёрнутый ответ.")
    document.add_paragraph("Ответ:")

    path = tmp_path / "ИСТ_ЕГЭ_Диагностика_21-22_Заданий 6.docx"
    document.save(path)
    return path


def test_converter_names_every_type_and_leaves_guesses_to_the_editor(tmp_path):
    result = tmp_path / "out.docx"
    outcome = converter.convert(_source(tmp_path), result, {})
    document = read_document(result)

    assert document.header == {
        "Предмет": "История", "Экзамен": "ЕГЭ", "Класс": "11", "Сезон": "21-22",
        "Тема": NEEDS_REVIEW,
    }
    by_number = {task.number: task for task in document.tasks}
    assert [task.number for task in document.tasks] == [1, 2, 3, 4, 5, 6]

    single = by_number[1]
    assert single.text("Тип") == "один ответ"
    assert single.text("Условие") == "Чему равна масса H_(2)O?"
    assert single.lines("Варианты") == ["1) 16 г", "2) 18 г"]
    assert single.text("Ответ") == "2"

    multiple = by_number[2]
    assert multiple.text("Тип") == "несколько ответов"
    assert multiple.lines("Варианты") == ["1) первое", "2) второе", "3) третье"]
    assert multiple.text("Ответ") == "1, 3"
    assert multiple.text("Решение") == "Верны первое и третье."

    matching = by_number[3]
    assert matching.text("Тип") == "соответствие"
    assert matching.lines("Пункты") == [
        "А) Восстание декабристов", "Б) Манифест о вольности", "В) Поход Игоря",
    ]
    assert matching.lines("Варианты") == ["1) 1185", "2) 1825", "3) 1762"]
    assert matching.text("Ответ") == "231"
    assert matching.text("Условие") == "Установите соответствие между событиями и годами."
    assert not any(block.kind == "table" for block in matching.fields["Условие"])
    assert outcome.dropped_hints == 1

    number = by_number[4]
    assert number.text("Тип") == "число"
    assert number.has_image("Условие")
    assert number.text("Ответ") == "1,5"
    assert number.text("Решение") == "Считаем."

    short = by_number[5]
    assert short.text("Тип") == "короткий ответ"
    assert short.lines("Ответ") == ["цитогенетический", "кариотипирование"]

    unread = by_number[6]
    assert unread.text("Тип") == NEEDS_REVIEW
    assert unread.text("Ответ") == NEEDS_REVIEW
    assert outcome.review_fields == [
        "Шапка, поле «Тема»", "Задание 6, поле «Тип»", "Задание 6, поле «Ответ»",
    ]

    problems = [str(problem) for problem in check_document(document)]
    assert problems == [
        "Шапка, поле «Тема»: помечено «ТРЕБУЕТ ПРОВЕРКИ».",
        "Задание 6, поле «Тип»: помечено «ТРЕБУЕТ ПРОВЕРКИ».",
        "Задание 6, поле «Ответ»: помечено «ТРЕБУЕТ ПРОВЕРКИ».",
    ]


def test_converter_carries_figures_as_parts_of_the_new_package(tmp_path):
    result = tmp_path / "out.docx"
    converter.convert(_source(tmp_path), result, {})
    reopened = Document(str(result))
    images = [
        rel for rel in reopened.part.rels.values() if rel.reltype.endswith("/image")
    ]
    assert len(images) == 1
    assert reopened.inline_shapes[0].width == Inches(1)


def test_topic_comes_from_the_filename_and_the_manifest_section():
    manifest = {
        "БИО_Цитология_Мейоз_Заданий 11.docx": {
            "exam": "ege", "grade": 11, "topic_hint": "11 класс/Задания/2. Цитология",
        },
    }
    header = converter.derive_header("БИО_Цитология_Мейоз_Заданий 11.docx", manifest)
    assert header == {
        "Предмет": "Биология", "Экзамен": "ЕГЭ", "Класс": "11",
        "Сезон": NEEDS_REVIEW, "Тема": "Мейоз",
    }
    header = converter.derive_header("МА_Угол_между_плоскостями.docx", {})
    assert header["Тема"] == "Угол между плоскостями"
    assert header["Экзамен"] == NEEDS_REVIEW
    header = converter.derive_header("ОБЩ 11_09.Налоги_ОДЗ_21-22_заданий 18.docx", {})
    assert header["Тема"] == "Налоги"
    assert header["Сезон"] == "21-22"


def _authoring(tmp_path: Path, answer: str, items: int = 3) -> Path:
    document = Document()
    for line in ("Предмет: История", "Экзамен: ЕГЭ", "Класс: 11", "Сезон: 21-22", "Тема: Даты"):
        document.add_paragraph(line)
    document.add_paragraph("Задание 1")
    document.add_paragraph("Тип: соответствие")
    document.add_paragraph("Условие:")
    document.add_paragraph("Установите соответствие.")
    document.add_paragraph("Пункты:")
    for label in "АБВГ"[:items]:
        document.add_paragraph(f"{label}) событие {label}")
    document.add_paragraph("Варианты:")
    for digit in "1234":
        document.add_paragraph(f"{digit}) год {digit}")
    document.add_paragraph(f"Ответ: {answer}")
    path = tmp_path / "authoring.docx"
    document.save(path)
    return path


def test_checker_names_the_task_and_field(tmp_path):
    clean = check_document(read_document(_authoring(tmp_path, "241")))
    assert clean == []

    problems = [str(problem) for problem in check_document(read_document(_authoring(tmp_path, "2413")))]
    assert problems == [
        "Задание 1, поле «Ответ»: у задания типа «соответствие» ответ должен состоять "
        "из 3 цифр по числу пунктов, получено 4."
    ]


def test_checker_rejects_figure_references_without_a_figure_and_form_hints(tmp_path):
    document = Document()
    for line in ("Предмет: Физика", "Экзамен: ЕГЭ", "Класс: 11", "Сезон: 21-22", "Тема: Оптика"):
        document.add_paragraph(line)
    document.add_paragraph("Задание 1")
    document.add_paragraph("Тип: число")
    document.add_paragraph("Условие:")
    document.add_paragraph("На рисунке показан ход луча.")
    document.add_paragraph("В ответ запишите только число.")
    document.add_paragraph("Ответ: 2")
    path = tmp_path / "figure.docx"
    document.save(path)
    problems = [str(problem) for problem in check_document(read_document(path))]
    assert problems == [
        "Задание 1, поле «Условие»: условие ссылается на рисунок, а рисунка нет.",
        "Задание 1, поле «Условие»: описание бланка ответа: «В ответ запишите только число.».",
    ]


def test_check_script_exits_nonzero_on_problems(tmp_path, capsys):
    from check_authoring_document import main

    assert main([str(_authoring(tmp_path, "241"))]) == 0
    assert capsys.readouterr().out.strip() == "ошибок нет"
    assert main([str(_authoring(tmp_path, "9"))]) == 1
    out = capsys.readouterr().out
    assert "Задание 1, поле «Ответ»" in out
