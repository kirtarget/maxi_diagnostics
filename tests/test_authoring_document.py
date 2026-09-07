"""The authoring format is a contract, so its reader is tested as one.

Every rule in `docs/AUTHORING_TEMPLATE.md` either holds here or the document is
refused with the task and the field named.
"""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
for entry in (str(ROOT), str(SCRIPTS)):
    if entry not in sys.path:
        sys.path.insert(0, entry)

import authoring_document as authoring  # noqa: E402
import convert_to_authoring_template as migrate  # noqa: E402
import make_authoring_example as example  # noqa: E402


HEADER = [
    "Предмет: Биология",
    "Экзамен: ЕГЭ",
    "Класс: 11",
    "Сезон: 21-22",
    "Тема: Цитология",
]
PICK_ONE = [
    "Задание 1",
    "Тип: один ответ",
    "Условие:",
    "Вопрос?",
    "Варианты:",
    "1) Первый",
    "2) Второй",
    "Ответ: 2",
]


def write(path: Path, lines: list[str]) -> Path:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(str(path))
    return path


def errors_of(tmp_path: Path, lines: list[str]) -> list[authoring.AuthoringError]:
    return authoring.check_authoring_document(write(tmp_path / "d.docx", lines))


def only_error(tmp_path: Path, lines: list[str]) -> authoring.AuthoringError:
    found = errors_of(tmp_path, lines)
    assert len(found) == 1, [str(error) for error in found]
    return found[0]


def test_the_example_document_reads_with_every_task_type(tmp_path):
    path = tmp_path / "example.docx"
    example.build(path)

    document = authoring.read_authoring_document(path)

    assert (document.subject, document.exam, document.grade) == ("Биология", "ЕГЭ", 11)
    assert document.season == "21-22"
    assert [task.kind for task in document.tasks] == [
        authoring.SINGLE, authoring.MULTIPLE, authoring.MATCHING,
        authoring.SEQUENCE, authoring.NUMBER, authoring.SHORT,
    ]

    pick_one, several, matching, sequence, number, short = document.tasks
    assert pick_one.options == ("Профаза", "Метафаза", "Анафаза", "Телофаза")
    assert pick_one.answers == ("2",)
    assert pick_one.explanation
    assert several.answers == ("1, 3",)
    assert matching.items[0] == "Цитогенетический"
    assert matching.answers == ("123",)
    # A task may name its own topic; the rest inherit the document's.
    assert matching.topic == "Методы биологических исследований"
    assert sequence.topic == number.topic == "Цитология"
    assert number.options == () and number.answers == ("46",)
    assert short.answers == ("цитогенетический", "кариотипирование", "цитологический")


def test_a_table_in_the_condition_stays_a_table(tmp_path):
    path = tmp_path / "example.docx"
    example.build(path)
    short = authoring.read_authoring_document(path).tasks[-1]

    assert len(short.prompt_tables) == 1
    assert short.prompt_tables[0].rows[0][0] == ("Метод",)
    assert all("|" not in block for block in short.prompt_blocks)


@pytest.mark.parametrize(
    ("line", "field"),
    [
        ("Предмет: Астрономия", "Предмет"),
        ("Экзамен: ВПР", "Экзамен"),
        ("Класс: одиннадцатый", "Класс"),
        ("Сезон: 2021", "Сезон"),
    ],
)
def test_the_header_refuses_a_value_it_does_not_know(tmp_path, line, field):
    header = [line if item.startswith(field) else item for item in HEADER]
    error = only_error(tmp_path, header + PICK_ONE)
    assert error.field == field
    assert error.task is None


def test_a_missing_header_field_is_named(tmp_path):
    error = only_error(tmp_path, HEADER[:-1] + PICK_ONE)
    assert error.field == "Тема"


def test_an_unknown_type_lists_the_ones_that_exist(tmp_path):
    lines = [line.replace("Тип: один ответ", "Тип: выбор") for line in PICK_ONE]
    error = only_error(tmp_path, HEADER + lines)
    assert (error.task, error.field) == (1, "Тип")
    assert authoring.MATCHING in str(error)


def test_a_pick_one_answer_must_name_an_option(tmp_path):
    lines = [line.replace("Ответ: 2", "Ответ: 5") for line in PICK_ONE]
    error = only_error(tmp_path, HEADER + lines)
    assert (error.task, error.field) == (1, "Ответ")


def test_a_matching_answer_is_as_long_as_its_items(tmp_path):
    task = [
        "Задание 1", "Тип: соответствие", "Условие:", "Соотнесите.",
        "Пункты:", "А) Первый", "Б) Второй",
        "Варианты:", "1) Раз", "2) Два",
        "Ответ: 121",
    ]
    error = only_error(tmp_path, HEADER + task)
    assert (error.task, error.field) == (1, "Ответ")
    assert "2 цифр" in str(error)


def test_a_sequence_may_repeat_a_digit_but_a_multiple_choice_may_not(tmp_path):
    sequence = [
        "Задание 1", "Тип: последовательность", "Условие:", "Заполните.",
        "Варианты:", "1) Раз", "2) Два", "Ответ: 122",
    ]
    assert errors_of(tmp_path, HEADER + sequence) == []

    several = [
        "Задание 1", "Тип: несколько ответов", "Условие:", "Выберите.",
        "Варианты:", "1) Раз", "2) Два", "Ответ: 1, 1",
    ]
    error = only_error(tmp_path, HEADER + several)
    assert (error.task, error.field) == (1, "Ответ")


def test_a_number_answer_must_be_a_number(tmp_path):
    task = ["Задание 1", "Тип: число", "Условие:", "Сколько?", "Ответ: сорок"]
    error = only_error(tmp_path, HEADER + task)
    assert (error.task, error.field) == (1, "Ответ")


def fields_of(tmp_path: Path, lines: list[str]) -> set[str]:
    return {error.field for error in errors_of(tmp_path, lines)}


def test_options_belong_to_the_types_that_have_them(tmp_path):
    without = ["Задание 1", "Тип: один ответ", "Условие:", "Вопрос?", "Ответ: 1"]
    assert "Варианты" in fields_of(tmp_path, HEADER + without)

    extra = [
        "Задание 1", "Тип: число", "Условие:", "Сколько?",
        "Варианты:", "1) Раз", "2) Два", "Ответ: 46",
    ]
    assert "Варианты" in fields_of(tmp_path, HEADER + extra)


def test_options_are_numbered_in_order_and_lose_their_labels(tmp_path):
    task = [
        "Задание 1", "Тип: один ответ", "Условие:", "Вопрос?",
        "Варианты:", "1) Первый", "3) Третий", "Ответ: 1",
    ]
    assert only_error(tmp_path, HEADER + task).field == "Варианты"

    document = write(tmp_path / "ok.docx", HEADER + PICK_ONE)
    assert authoring.read_authoring_document(document).tasks[0].options == ("Первый", "Второй")


def test_a_field_cannot_appear_twice(tmp_path):
    found = errors_of(tmp_path, HEADER + PICK_ONE + ["Тип: число"])
    assert found and found[0].task == 1 and found[0].field == "Тип"
    assert "второй раз" in str(found[0])


def test_a_line_outside_a_field_is_refused(tmp_path):
    lines = ["Задание 1", "Просто текст", "Тип: один ответ"]
    error = only_error(tmp_path, HEADER + lines)
    assert error.task == 1


def test_tasks_are_numbered_in_order(tmp_path):
    lines = [line.replace("Задание 1", "Задание 2") for line in PICK_ONE]
    error = only_error(tmp_path, HEADER + lines)
    assert error.task == 2


def test_a_converted_editorial_document_marks_what_it_could_not_read(tmp_path):
    """The migration fills what is certain and hands the rest to the editor."""
    legacy = tmp_path / "legacy.docx"
    document = Document()
    document.add_paragraph("Задание 1")
    document.add_paragraph("В какой фазе хромосомы на экваторе?")
    document.add_paragraph("Варианты:")
    document.add_paragraph("Профаза")
    document.add_paragraph("Метафаза")
    document.add_paragraph("Ответ:")
    document.add_paragraph("2")
    document.add_paragraph("Задание 2")
    document.add_paragraph("Установите последовательность.")
    document.add_paragraph("Ответ:")
    document.add_paragraph("231")
    document.save(str(legacy))

    report = migrate.convert(legacy, tmp_path / "authoring.docx")

    assert report["tasks"] == 2
    assert report["clean"] == 1
    assert [number for number, _ in report["marked"]] == [2]
    assert "Тема" in report["header"]

    text = "\n".join(
        paragraph.text for paragraph in Document(str(tmp_path / "authoring.docx")).paragraphs
    )
    assert "Тип: один ответ" in text
    assert f"Тип: {authoring.SEQUENCE}" in text
    assert migrate.REVIEW in text
